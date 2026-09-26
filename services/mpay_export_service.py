"""
M-Pay Export Service (services/mpay_export_service.py)
Geração de planilhas formatadas (.xlsx) e arquivos .csv para o módulo M-Pay.
"""

from __future__ import annotations

import csv
import io
from datetime import datetime
from typing import Any

from flask import Response, send_file


def build_mpay_xlsx(rows: list[dict[str, Any]]) -> Response:
    """Gera uma planilha Excel estilizada profissionalmente via openpyxl."""
    import openpyxl
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "M-Pay Comprovantes"

    # Título do Relatório
    ws.merge_cells("A1:J1")
    title_cell = ws["A1"]
    title_cell.value = "M-PAY • RELATÓRIO DE COMPROVANTES DE PAGAMENTO"
    title_cell.font = Font(size=14, bold=True, color="1E3A8A")
    title_cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 32

    headers = [
        "Data Pagamento",
        "Empresa Pagadora",
        "Forma / Origem",
        "Beneficiário / Destino",
        "Documento / Chave",
        "Valor (R$)",
        "Banco Origem",
        "Método",
        "Categoria",
        "Observações",
    ]
    ws.append(headers)

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="2563EB", end_color="2563EB", fill_type="solid")
    border_thin = Border(
        left=Side(style="thin", color="CCCCCC"),
        right=Side(style="thin", color="CCCCCC"),
        top=Side(style="thin", color="CCCCCC"),
        bottom=Side(style="thin", color="CCCCCC"),
    )

    ws.row_dimensions[2].height = 24
    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=2, column=col_idx)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")

    total_amount = 0.0
    for r in rows:
        val = float(r.get("amount") or 0.0)
        total_amount += val
        pdate = r.get("paid_at").strftime("%d/%m/%Y") if r.get("paid_at") else ""
        row_vals = [
            pdate,
            r.get("paying_company") or "M-one",
            r.get("payment_source") or "Conta da Empresa",
            r.get("beneficiary_name") or "",
            r.get("beneficiary_document") or "",
            val,
            r.get("bank_origin") or "",
            r.get("payment_method") or "",
            r.get("category") or "",
            r.get("notes") or "",
        ]
        ws.append(row_vals)
        current_row = ws.max_row
        ws.cell(row=current_row, column=6).number_format = '"R$ "#,##0.00'
        for c in range(1, len(headers) + 1):
            ws.cell(row=current_row, column=c).border = border_thin

    # Linha de Total
    total_row = ws.max_row + 1
    ws.cell(row=total_row, column=5, value="TOTAL GERAL:").font = Font(bold=True)
    total_cell = ws.cell(row=total_row, column=6, value=total_amount)
    total_cell.font = Font(bold=True, color="047857")
    total_cell.number_format = '"R$ "#,##0.00'

    # Ajuste de larguras das colunas
    col_widths = [16, 20, 20, 32, 22, 18, 18, 14, 18, 28]
    for i, w in enumerate(col_widths, 1):
        col_letter = openpyxl.utils.get_column_letter(i)
        ws.column_dimensions[col_letter].width = w

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    filename = f"mpay_comprovantes_{datetime.now().strftime('%Y%m%d_%H%M')}.xlsx"
    return send_file(
        output,
        download_name=filename,
        as_attachment=True,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


def build_mpay_csv(rows: list[dict[str, Any]]) -> Response:
    """Gera exportação em CSV com UTF-8 BOM e delimitador ponto-e-vírgula."""
    output = io.StringIO()
    output.write("\ufeff")  # BOM para Excel
    writer = csv.writer(output, delimiter=";")
    writer.writerow(["Data", "Empresa Pagadora", "Forma Pagamento", "Favorecido", "CPF/CNPJ/Pix", "Valor R$", "Banco", "Metodo", "Categoria", "Observacoes"])

    for r in rows:
        pdate = r.get("paid_at").strftime("%d/%m/%Y") if r.get("paid_at") else ""
        val = f"{float(r.get('amount') or 0.0):.2f}".replace(".", ",")
        writer.writerow([
            pdate,
            r.get("paying_company") or "M-one",
            r.get("payment_source") or "Conta da Empresa",
            r.get("beneficiary_name") or "",
            r.get("beneficiary_document") or "",
            val,
            r.get("bank_origin") or "",
            r.get("payment_method") or "",
            r.get("category") or "",
            r.get("notes") or "",
        ])

    csv_data = output.getvalue()
    filename = f"mpay_comprovantes_{datetime.now().strftime('%Y%m%d_%H%M')}.csv"
    return Response(
        csv_data,
        mimetype="text/csv; charset=utf-8",
        headers={"Content-Disposition": f"attachment; filename={filename}"},
    )


def export_mpay_dataset(export_format: str, month_filter: str = "", search_q: str = "") -> Response:
    """Filtra os dados e despacha a exportação nos formatos XLSX ou CSV."""
    from database import db

    sql_where = []
    params: list[Any] = []
    if month_filter:
        sql_where.append("TO_CHAR(paid_at, 'YYYY-MM') = %s")
        params.append(month_filter)
    if search_q:
        sql_where.append("(beneficiary_name ILIKE %s OR notes ILIKE %s OR bank_origin ILIKE %s OR paying_company ILIKE %s OR payment_source ILIKE %s)")
        params.extend([f"%{search_q}%"] * 5)

    where_clause = f"WHERE {' AND '.join(sql_where)}" if sql_where else ""

    with db() as conn:
        rows = conn.execute(
            f"""
            SELECT paid_at, paying_company, payment_source, beneficiary_name, beneficiary_document, amount,
                   bank_origin, payment_method, category, notes, confidence_status
            FROM mpay_transactions
            {where_clause}
            ORDER BY paid_at DESC NULLS LAST, id DESC
            """,
            params,
        ).fetchall()

    if export_format == "xlsx":
        return build_mpay_xlsx(rows)
    return build_mpay_csv(rows)


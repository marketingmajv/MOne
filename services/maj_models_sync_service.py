"""
Serviço de Sincronização 24/7 dos Modelos MAJ Mobilidade
Lê as 9 abas estruturadas da Planilha Google oficial de produtos,
cruza especificações técnicas (Desempenho, Bateria, Estrutura, Segurança, Comercial, Texto Livre)
e gera o catálogo canônico unificado para o Mostruário /modelos e a API /api/models.
Permite ativar/inativar modelos de forma persistente.
"""

import json
import logging
import re
import time
import urllib.request
import io
from pathlib import Path
from typing import Dict, List, Optional, Any, Set

logger = logging.getLogger(__name__)

BASE_DIR = Path(__file__).resolve().parent.parent
CACHE_FILE = BASE_DIR / "static" / "data" / "maj_models.json"
API_CACHE_FILE = BASE_DIR / "api" / "static" / "data" / "maj_models.json"
STATUS_FILE = BASE_DIR / "static" / "data" / "model_status.json"
API_STATUS_FILE = BASE_DIR / "api" / "static" / "data" / "model_status.json"
ORDER_FILE = BASE_DIR / "static" / "data" / "models_order.json"
API_ORDER_FILE = BASE_DIR / "api" / "static" / "data" / "models_order.json"

SHEET_URL = (
    "https://docs.google.com/spreadsheets/d/"
    "1uYnE9-MCSuRSe-9bSvcUulALjvh8aUdYEMbGAWZEC5E/export?format=xlsx"
)

# Mapeamento oficial de slugs para diretórios de fotos existentes em public/images/
SLUG_ALIASES = {
    "MAX 12 MAJ": "max-12",
    "VITTÓRIA": "vittoria",
    "VITTORIA": "vittoria",
    "SPORT 701 SHORT": "sport-701-short",
    "CLASSIC 1000": "classic-1000",
    "FLOW ONE": "flow-one",
    "NOVA": "nova",
    "RZ 110": "rz-110",
    "RIDE ON": "ride-on",
    "M9 PRO": "m9-pro",
    "V80 PRO": "v80-pro",
    "M50 PRO": "m50-pro",
    "V20 ULTRA": "v20-ultra",
    "V8 MINI ULTRA": "v8-mini",
    "V8 MINI": "v8-mini",
    "H1": "h1",
}

# Modelos inativos por padrão (ex: H1 removido do catálogo público)
DEFAULT_INACTIVE_SLUGS: Set[str] = {"h1"}

_memory_cache: Optional[Dict[str, Any]] = None
_last_sync_time: float = 0
CACHE_TTL_SECONDS = 3600  # 1 hora de cache em memória


def load_inactive_slugs() -> Set[str]:
    """Carrega lista de slugs inativados a partir do arquivo persistente de status."""
    inactive = set(DEFAULT_INACTIVE_SLUGS)
    for p in [STATUS_FILE, API_STATUS_FILE]:
        if p.is_file():
            try:
                with open(p, "r", encoding="utf-8") as f:
                    data = json.load(f)
                    for slug, is_active in data.items():
                        if not is_active:
                            inactive.add(slug.lower().strip())
                        elif slug.lower().strip() in inactive:
                            inactive.remove(slug.lower().strip())
                break
            except Exception as e:
                logger.warning(f"Erro ao ler {p}: {e}")
    return inactive


def set_model_active_status(slug: str, is_active: bool) -> bool:
    """Ativa ou inativa um modelo específico no backend e persiste a configuração."""
    global _memory_cache, _last_sync_time
    clean_slug = slug.strip().lower()

    current_status = {}
    for p in [STATUS_FILE, API_STATUS_FILE]:
        if p.is_file():
            try:
                with open(p, "r", encoding="utf-8") as f:
                    current_status = json.load(f)
                break
            except Exception:
                pass

    current_status[clean_slug] = bool(is_active)

    for p in [STATUS_FILE, API_STATUS_FILE]:
        try:
            p.parent.mkdir(parents=True, exist_ok=True)
            with open(p, "w", encoding="utf-8") as f:
                json.dump(current_status, f, ensure_ascii=False, indent=2)
        except Exception as e:
            logger.error(f"Erro ao salvar status em {p}: {e}")

    # Forçar recálculo do cache
    _memory_cache = None
    _last_sync_time = 0
    sync_models_from_sheets(force=True)
    return True


def load_models_order() -> List[str]:
    """Carrega a ordem customizada dos modelos a partir do arquivo persistente."""
    for p in [ORDER_FILE, API_ORDER_FILE]:
        if p.is_file():
            try:
                with open(p, "r", encoding="utf-8") as f:
                    data = json.load(f)
                    if isinstance(data, list):
                        return [str(s).strip().lower() for s in data if str(s).strip()]
            except Exception as e:
                logger.warning(f"Erro ao ler {p}: {e}")
    return []


def save_models_order(order_slugs: List[str]) -> bool:
    """Salva a ordem customizada dos modelos e sincroniza o cache em memória."""
    global _memory_cache, _last_sync_time
    clean_order = [str(s).strip().lower() for s in order_slugs if str(s).strip()]
    for p in [ORDER_FILE, API_ORDER_FILE]:
        try:
            p.parent.mkdir(parents=True, exist_ok=True)
            with open(p, "w", encoding="utf-8") as f:
                json.dump(clean_order, f, ensure_ascii=False, indent=2)
        except Exception as e:
            logger.error(f"Erro ao salvar ordem em {p}: {e}")
            return False
    _memory_cache = None
    _last_sync_time = 0
    sync_models_from_sheets(force=True)
    return True


def normalize_slug(name: str) -> str:
    """Gera slug padronizado e resolve contra aliases conhecidos."""
    clean_name = name.strip().upper()
    if clean_name in SLUG_ALIASES:
        return SLUG_ALIASES[clean_name]

    slug = name.lower().strip()
    slug = re.sub(r"[àáâãäå]", "a", slug)
    slug = re.sub(r"[èéêë]", "e", slug)
    slug = re.sub(r"[ìíîï]", "i", slug)
    slug = re.sub(r"[òóôõö]", "o", slug)
    slug = re.sub(r"[ùúûü]", "u", slug)
    slug = re.sub(r"[ç]", "c", slug)
    slug = re.sub(r"[^a-z0-9]+", "-", slug).strip("-")
    return slug


def find_model_images(slug: str) -> List[str]:
    """Descobre fotos existentes locais em public/images/{slug}/."""
    images = []
    search_dirs = [
        BASE_DIR / "public" / "images" / slug,
        BASE_DIR / "static" / "images" / slug,
        BASE_DIR / "public" / "images" / slug.replace("-", ""),
    ]
    for d in search_dirs:
        if d.is_dir():
            for f in sorted(d.glob("*.jpg")) + sorted(d.glob("*.png")) + sorted(d.glob("*.webp")):
                if f.name.startswith("."):
                    continue
                images.append(f"/static/images/{slug}/{f.name}" if "static" in str(d) else f"/public/images/{slug}/{f.name}")
            if images:
                break
    return images


def sync_models_from_sheets(force: bool = False) -> Dict[str, Any]:
    """Baixa o XLSX oficial das 9 abas do Google Sheets e compila o catálogo canônico."""
    global _memory_cache, _last_sync_time
    now = time.time()

    if not force and _memory_cache and (now - _last_sync_time < CACHE_TTL_SECONDS):
        return _memory_cache

    try:
        import openpyxl

        req = urllib.request.Request(SHEET_URL, headers={"User-Agent": "MOne-SyncEngine/1.0"})
        with urllib.request.urlopen(req, timeout=15) as resp:
            data = resp.read()

        wb = openpyxl.load_workbook(io.BytesIO(data), data_only=True)
        models_dict: Dict[str, Dict[str, Any]] = {}
        inactive_slugs = load_inactive_slugs()

        for sheet in wb.sheetnames:
            if sheet in ["Instruções", "Resumo"]:
                continue
            ws = wb[sheet]
            rows = list(ws.iter_rows(values_only=True))
            if len(rows) < 4:
                continue

            headers = [str(c).strip() if c is not None else f"col_{i}" for i, c in enumerate(rows[2])]

            model_col_idx = 2 if "Marca" in headers and headers.index("Marca") < headers.index("Modelo") else 1
            cat_col_idx = headers.index("Categoria") if "Categoria" in headers else 0

            for r in rows[3:]:
                if not r or len(r) <= model_col_idx or not r[model_col_idx]:
                    continue
                name = str(r[model_col_idx]).strip()
                cat = str(r[cat_col_idx]).strip() if len(r) > cat_col_idx and r[cat_col_idx] else ""
                if not name or name.lower() in ["modelo", "categoria", "none"]:
                    continue

                m_key = name.upper()
                if m_key not in models_dict:
                    slug = normalize_slug(name)
                    cutout_file = BASE_DIR / "public" / "images" / "models" / "cutouts" / f"{slug}.png"
                    cutout_url = f"/public/images/models/cutouts/{slug}.png" if cutout_file.exists() else ""
                    images = find_model_images(slug)
                    hero_url = cutout_url if cutout_url else (images[0] if images else "")

                    models_dict[m_key] = {
                        "categoria": cat,
                        "modelo": name,
                        "slug": slug,
                        "is_active": slug not in inactive_slugs,
                        "sections": {},
                        "cutout_image": cutout_url,
                        "hero_image": hero_url,
                        "images": images,
                    }
                elif cat and not models_dict[m_key]["categoria"]:
                    models_dict[m_key]["categoria"] = cat

                sheet_data = {}
                for h, val in zip(headers, r):
                    # Remove referência de potência de pico
                    if h in ["Categoria", "Modelo", "% preenchido", "Potência de pico (W)"]:
                        continue
                    if val is not None:
                        val_str = str(val).strip()
                        if val_str.endswith(".0") and val_str[:-2].isdigit():
                            val_str = val_str[:-2]
                        sheet_data[h] = val_str
                    else:
                        sheet_data[h] = ""

                models_dict[m_key]["sections"][sheet] = sheet_data

        # Enriquecer cada modelo com as métricas canônicas (Apenas Potência Nominal)
        final_list = []
        for m_key, m_val in models_dict.items():
            desemp = m_val["sections"].get("Desempenho", {})
            bateria = m_val["sections"].get("Bateria e autonomia", {})
            estrut = m_val["sections"].get("Estrutura e conforto", {})
            comercial = m_val["sections"].get("Uso e comercial", {})
            livre = m_val["sections"].get("Texto livre", {})
            seguranca = m_val["sections"].get("Segurança e componentes", {})

            pot_nom = desemp.get("Potência nominal (W)", "")
            vel_max = desemp.get("Velocidade máxima (km/h)", "")
            autonomia_cons = bateria.get("Autonomia conservadora (km)", "")
            autonomia_max = bateria.get("Autonomia máxima estimada (km)", "")

            # Métricas sem referência a pico de potência
            m_val["metrics"] = {
                "potencia_nominal_w": pot_nom,
                "velocidade_max_kmh": vel_max,
                "autonomia_conservadora_km": autonomia_cons,
                "autonomia_maxima_km": autonomia_max,
                "voltagem_v": bateria.get("Voltagem (V)", ""),
                "capacidade_ah": bateria.get("Capacidade (Ah)", ""),
                "capacidade_wh": bateria.get("Capacidade (Wh)", ""),
                "quimica_bateria": bateria.get("Química / tipo", ""),
                "carga_maxima_kg": estrut.get("Carga máxima (kg)", ""),
                "peso_kg": estrut.get("Peso do veículo (kg)", ""),
                "freios": f"Diant: {seguranca.get('Freio dianteiro', '')} / Tras: {seguranca.get('Freio traseiro', '')}".strip(" /"),
                "diferencial": comercial.get("Principal diferencial", ""),
                "descricao_livre": livre.get("Descrição comercial livre", ""),
                "pontos_fortes": livre.get("Pontos fortes", ""),
            }
            final_list.append(m_val)

        # Ordenar por Ordem Customizada (se definida) ou por Categoria/Nome
        custom_order = load_models_order()
        if custom_order:
            order_map = {slug: idx for idx, slug in enumerate(custom_order)}
            category_order = {"Scooter": 1, "Bike": 2, "Triciclo": 3}
            final_list.sort(key=lambda x: (
                order_map.get(x["slug"], 900 + category_order.get(x["categoria"], 99)),
                x["modelo"]
            ))
        else:
            category_order = {"Scooter": 1, "Bike": 2, "Triciclo": 3}
            final_list.sort(key=lambda x: (category_order.get(x["categoria"], 99), x["modelo"]))

        active_models = [m for m in final_list if m["is_active"]]

        result = {
            "updated_at": time.strftime("%Y-%m-%dT%H:%M:%SZ", time.gmtime()),
            "total_models": len(active_models),
            "total_all": len(final_list),
            "categories": list(dict.fromkeys(m["categoria"] for m in active_models if m["categoria"])),
            "models": active_models,
            "all_models": final_list,
        }

        # Salvar em cache no disco para persistência
        for cache_path in [CACHE_FILE, API_CACHE_FILE]:
            try:
                cache_path.parent.mkdir(parents=True, exist_ok=True)
                with open(cache_path, "w", encoding="utf-8") as f:
                    json.dump(result, f, ensure_ascii=False, indent=2)
            except Exception as ce:
                logger.warning(f"Erro ao salvar cache em {cache_path}: {ce}")

        _memory_cache = result
        _last_sync_time = now
        logger.info(f"Sincronização MAJ Models concluída: {len(active_models)} ativos ({len(final_list)} totais).")
        return result

    except Exception as e:
        logger.error(f"Falha ao sincronizar com Google Sheets: {e}", exc_info=True)
        if _memory_cache:
            return _memory_cache
        for cache_path in [CACHE_FILE, API_CACHE_FILE]:
            if cache_path.is_file():
                try:
                    with open(cache_path, "r", encoding="utf-8") as f:
                        disk_data = json.load(f)
                    _memory_cache = disk_data
                    return disk_data
                except Exception:
                    pass
        return {"error": str(e), "total_models": 0, "models": []}


def get_all_models(include_inactive: bool = False) -> List[Dict[str, Any]]:
    """Retorna os modelos sincronizados (ativos por padrão)."""
    data = sync_models_from_sheets(force=False)
    if include_inactive:
        return data.get("all_models", data.get("models", []))
    return data.get("models", [])


def get_model_by_slug(slug: str) -> Optional[Dict[str, Any]]:
    """Busca um modelo específico por slug."""
    models = get_all_models(include_inactive=True)
    clean = slug.strip().lower()
    for m in models:
        if m.get("slug") == clean:
            return m
    return None

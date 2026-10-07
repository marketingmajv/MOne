"""
M-One Autonomous Freight AI Pipeline (services/freight_ai_pipeline.py)
Motor de IA em 2 estágios para descoberta e compilação autônoma de qualquer formato de frete.
1. Estágio 1: Inspeção de layout e classificação semântica de abas/seções via Gemini.
2. Estágio 2: Normalização com Forward-Fill e extração estruturada de tarifas e SLA por cidade.
"""

from __future__ import annotations

import json
import logging
import os
import re
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple

import openpyxl

from services.gemini_client import execute_gemini_payload, get_gemini_api_key

logger = logging.getLogger("mone.freight_ai")

BRAZIL_UFS = {
    "AC", "AL", "AP", "AM", "BA", "CE", "DF", "ES", "GO", "MA",
    "MT", "MS", "MG", "PA", "PB", "PR", "PE", "PI", "RJ", "RN",
    "RS", "RO", "RR", "SC", "SP", "SE", "TO"
}

DEFAULT_REGIONAL_DAYS = {
    "ES": 2, "RJ": 3, "SP": 3, "MG": 4, "BA": 5, "SE": 5, "AL": 6, "PE": 6,
    "PB": 6, "RN": 7, "CE": 7, "PI": 8, "MA": 8, "PA": 8, "TO": 7, "GO": 5,
    "DF": 4, "PR": 4, "SC": 5, "RS": 5, "MT": 7, "MS": 6, "RO": 9, "AC": 10,
    "AM": 10, "RR": 12, "AP": 10
}


def clean_num(val: Any, default: float = 0.0) -> float:
    """Converte números com vírgula ou pontuação em float seguro."""
    if val is None:
        return default
    if isinstance(val, (int, float)):
        return float(val)
    s = str(val).strip()
    if not s or any(w in s.upper() for w in ["SOB", "CONSULTA", "N/A", "-"]):
        return default
    s = re.sub(r"[^\d,\.-]", "", s)
    if "," in s and "." in s:
        s = s.replace(".", "").replace(",", ".")
    elif "," in s:
        s = s.replace(",", ".")
    try:
        return float(s)
    except (ValueError, TypeError):
        return default


def inspect_workbook_layout_with_ai(file_path: str) -> Dict[str, Any]:
    """
    Estágio 1: Analisa as abas e cabeçalhos da planilha com Gemini e classifica
    autonomamente qual aba traz as tarifas e qual traz os prazos/cidades.
    """
    p = Path(file_path)
    if p.suffix.lower() not in [".xlsx", ".xlsm", ".xltx", ".xls"]:
        return {}

    try:
        wb = openpyxl.load_workbook(file_path, data_only=True, read_only=True)
    except Exception as e:
        logger.warning("Falha ao abrir planilha com openpyxl: %s", e)
        return {}

    catalog = []
    for sheet_name in wb.sheetnames:
        sheet = wb[sheet_name]
        sample_rows = []
        for r_idx, r in enumerate(sheet.iter_rows(values_only=True)):
            if r_idx > 12:
                break
            cells = [str(c).strip() for c in r if c is not None and str(c).strip() != ""]
            if cells:
                sample_rows.append(" | ".join(cells[:10]))
        catalog.append({"sheet_name": sheet_name, "sample": "\n".join(sample_rows[:5])})

    prompt = f"""
    Você é um especialista em logística e fretes rodoviários no Brasil.
    Analise o catálogo de abas da planilha enviada e identifique:
    1. 'pricing_sheet': Qual aba contém as TARIFAS / PREÇOS (faixas de peso em kg, valores em R$, capitais, interior).
    2. 'sla_sheet': Qual aba contém a RELAÇÃO DE CIDADES / PRAZOS DE ENTREGA (SLA em dias úteis).
    3. 'carrier_name': Nome da transportadora (se mencionado no cabeçalho, notas ou e-mail de contato, ex: 'Unitrans Transportes', 'Águia Branca').

    CATÁLOGO DE ABAS:
    {json.dumps(catalog, ensure_ascii=False, indent=2)}

    Retorne ESTRITAMENTE um JSON:
    ```json
    {{
      "pricing_sheet": "nome da aba de tarifas",
      "sla_sheet": "nome da aba de cidades/prazos ou null",
      "carrier_name": "Nome da Transportadora ou null"
    }}
    ```
    """

    res = execute_gemini_payload({
        "contents": [{"parts": [{"text": prompt}]}],
        "generationConfig": {"temperature": 0.1, "maxOutputTokens": 1024}
    })

    pricing_sheet = None
    sla_sheet = None
    carrier_name = None

    if res.get("success"):
        try:
            raw = res.get("text", "")
            m = re.search(r"```json\s*(.*?)\s*```", raw, re.DOTALL)
            data = json.loads(m.group(1) if m else raw)
            pricing_sheet = data.get("pricing_sheet")
            sla_sheet = data.get("sla_sheet")
            carrier_name = data.get("carrier_name")
        except Exception:
            pass

    # Heurística de contingência resiliente
    if not pricing_sheet or pricing_sheet not in wb.sheetnames:
        for item in catalog:
            s_name = item["sheet_name"].upper()
            s_text = (item["sheet_name"] + " " + item["sample"]).upper()
            if any(w in s_name for w in ["TABELA", "PRECO", "TARIFA"]) and any(w in s_text for w in ["KG", "50", "100", "EXCEDENTE", "DESTINO"]):
                pricing_sheet = item["sheet_name"]
                break
        if not pricing_sheet and wb.sheetnames:
            pricing_sheet = wb.sheetnames[0]

    if not sla_sheet or sla_sheet not in wb.sheetnames:
        for item in catalog:
            s_name = item["sheet_name"].upper()
            if any(w in s_name for w in ["CIDADE", "PRAZO", "SLA", "TDA"]) and item["sheet_name"] != pricing_sheet:
                sla_sheet = item["sheet_name"]
                break

    # Detecção de nome da transportadora no texto
    if not carrier_name:
        all_text = " ".join([c["sample"] for c in catalog])
        if "UNITRANS" in all_text.upper():
            carrier_name = "Unitrans Transportes"
        elif "AGUIA BRANCA" in all_text.upper():
            carrier_name = "Viação Águia Branca S/A"

    return {
        "pricing_sheet": pricing_sheet,
        "sla_sheet": sla_sheet,
        "carrier_name": carrier_name or "Transportadora"
    }


def parse_pricing_sheet_matrix(sheet: openpyxl.worksheet.worksheet.Worksheet) -> List[Dict[str, Any]]:
    """
    Extrai as matrizes tarifárias com Forward-Fill inteligente de UF e detecção flexível de colunas.
    Suporta faixas de 0-50kg, 100kg, excedente, pedágio, ad-valorem, despacho e prazos gerais.
    """
    rows = list(sheet.iter_rows(values_only=True))
    if len(rows) < 3:
        return []

    # 1. Localizar linha de cabeçalho
    header_idx = -1
    col_map: Dict[str, int] = {}

    for idx, r in enumerate(rows[:15]):
        str_cells = [str(c or "").strip().upper() for c in r]
        if any("UF" in c or "ESTADO" in c for c in str_cells) and any("50" in c or "100" in c or "PESO" in c or "KG" in c or "VALOR" in c for c in str_cells):
            header_idx = idx
            for c_i, c_val in enumerate(str_cells):
                if c_val in ["UF", "ESTADO"]: col_map["uf"] = c_i
                elif any(w in c_val for w in ["DESTINO", "TIPO", "PRACA"]): col_map["dest"] = c_i
                elif "50" in c_val and ("ATE" in c_val or "KG" in c_val or c_val == "50KG"): col_map["p50"] = c_i
                elif "100" in c_val: col_map["p100"] = c_i
                elif any(w in c_val for w in ["EXCED", "SOBRE", "KG EXC"]): col_map["exced"] = c_i
                elif any(w in c_val for w in ["PEDAG", "PED"]): col_map["pedagio"] = c_i
                elif any(w in c_val for w in ["AD", "GRIS", "SEGURO"]): col_map["ad_val"] = c_i
                elif any(w in c_val for w in ["DESPACHO", "TAXA"]): col_map["despacho"] = c_i
                elif any(w in c_val for w in ["PRAZO", "DIAS"]): col_map["prazo"] = c_i
            break

    if header_idx == -1 or "uf" not in col_map:
        return []

    current_uf: Optional[str] = None
    master_rates: List[Dict[str, Any]] = []

    for r in rows[header_idx + 1:]:
        if not any(r):
            continue
        line_str = " ".join([str(c or "") for c in r]).upper()
        if "CONSULTA" in line_str or "GENERALIDADES" in line_str:
            continue

        uf_raw = str(r[col_map["uf"]] or "").strip().upper()
        if uf_raw in BRAZIL_UFS:
            current_uf = uf_raw

        dest_raw = str(r[col_map.get("dest", 1)] or "").strip().upper() if "dest" in col_map else "GERAL"
        if not current_uf or not dest_raw:
            continue

        p50 = clean_num(r[col_map["p50"]]) if "p50" in col_map and col_map["p50"] < len(r) else 0.0
        p100 = clean_num(r[col_map["p100"]]) if "p100" in col_map and col_map["p100"] < len(r) else 0.0
        exced = clean_num(r[col_map["exced"]]) if "exced" in col_map and col_map["exced"] < len(r) else 0.0
        ped = clean_num(r[col_map["pedagio"]]) if "pedagio" in col_map and col_map["pedagio"] < len(r) else 0.0
        ad_v = clean_num(r[col_map["ad_val"]]) if "ad_val" in col_map and col_map["ad_val"] < len(r) else 0.0
        ad_v_pct = round(ad_v * 100, 3) if 0 < ad_v < 0.1 else ad_v
        desp = clean_num(r[col_map["despacho"]]) if "despacho" in col_map and col_map["despacho"] < len(r) else 0.0
        
        prazo_raw = str(r[col_map["prazo"]] or "") if "prazo" in col_map and col_map["prazo"] < len(r) else ""
        m_days = re.search(r"(\d+)", prazo_raw)
        days = int(m_days.group(1)) if m_days else DEFAULT_REGIONAL_DAYS.get(current_uf, 5)

        if p50 > 0 or p100 > 0:
            fixed_base = p50 if p50 > 0 else p100
            p100_val = p100 if p100 > 0 else p50

            # Faixa 1: Até 50kg
            master_rates.append({
                "uf": current_uf,
                "dest_type": dest_raw,
                "city": None,
                "min_weight": 0.0,
                "max_weight": 50.0,
                "fixed_price": round(fixed_base + desp, 2),
                "weight_price_per_kg": 0.0,
                "ad_valorem_percent": ad_v_pct,
                "toll_per_100kg": ped,
                "dispatch_fixed": desp,
                "delivery_days": days,
                "notes": f"{dest_raw} (Até 50kg) • Pedágio R${ped:.2f}/100kg"
            })

            # Faixa 2: 50.01 a 100kg
            master_rates.append({
                "uf": current_uf,
                "dest_type": dest_raw,
                "city": None,
                "min_weight": 50.01,
                "max_weight": 100.0,
                "fixed_price": round(p100_val + desp, 2),
                "weight_price_per_kg": 0.0,
                "ad_valorem_percent": ad_v_pct,
                "toll_per_100kg": ped,
                "dispatch_fixed": desp,
                "delivery_days": days,
                "notes": f"{dest_raw} (Até 100kg) • Pedágio R${ped:.2f}/100kg"
            })

            # Faixa 3: Excedente > 100kg
            master_rates.append({
                "uf": current_uf,
                "dest_type": dest_raw,
                "city": None,
                "min_weight": 100.01,
                "max_weight": 999999.0,
                "fixed_price": round(p100_val + desp, 2),
                "weight_price_per_kg": exced,
                "ad_valorem_percent": ad_v_pct,
                "toll_per_100kg": ped,
                "dispatch_fixed": desp,
                "delivery_days": days,
                "notes": f"{dest_raw} Excedente R${exced:.2f}/kg • Pedágio R${ped:.2f}/100kg"
            })

    return master_rates


def parse_sla_sheet_cities(sheet: openpyxl.worksheet.worksheet.Worksheet) -> Dict[str, Dict[str, Dict[str, Any]]]:
    """Extrai cidades, prazos em dias úteis e TDA da aba de SLA."""
    rows = list(sheet.iter_rows(values_only=True))
    if len(rows) < 2:
        return {}

    uf_col = -1
    city_col = -1
    desc_col = -1
    prazo_col = -1
    tda_col = -1

    for idx, r in enumerate(rows[:10]):
        str_cells = [str(c or "").strip().upper() for c in r]
        if any("UF" in c for c in str_cells) and any("CIDADE" in c or "MUNIC" in c for c in str_cells):
            for c_i, c_val in enumerate(str_cells):
                if c_val == "UF": uf_col = c_i
                elif any(w in c_val for w in ["CIDADE", "MUNIC"]): city_col = c_i
                elif any(w in c_val for w in ["DESCRI", "TIPO"]): desc_col = c_i
                elif "PRAZO" in c_val: prazo_col = c_i
                elif "TDA" in c_val: tda_col = c_i
            break

    if uf_col == -1 or city_col == -1 or prazo_col == -1:
        return {}

    cities_map: Dict[str, Dict[str, Dict[str, Any]]] = {}
    for r in rows[1:]:
        if not any(r): continue
        uf_val = str(r[uf_col] or "").strip().upper()
        city_val = str(r[city_col] or "").strip().upper()
        if uf_val in BRAZIL_UFS and city_val:
            prazo_raw = str(r[prazo_col] or "").strip()
            m = re.search(r"(\d+)", prazo_raw)
            days = int(m.group(1)) if m else DEFAULT_REGIONAL_DAYS.get(uf_val, 5)
            tda_val = clean_num(r[tda_col]) if tda_col != -1 and tda_col < len(r) else 0.0
            desc_val = str(r[desc_col] or "").strip().upper() if desc_col != -1 and desc_col < len(r) else "INTERIOR"

            cities_map.setdefault(uf_val, {})[city_val] = {
                "days": days,
                "tda": tda_val,
                "dest_type": desc_val
            }

    return cities_map


def run_autonomous_freight_pipeline(
    file_path: str,
    carrier_hint: str = "",
    table_hint: str = ""
) -> Dict[str, Any]:
    """Pipeline Autônomo Unificado (All-in-One e Multi-Aba)."""
    layout = inspect_workbook_layout_with_ai(file_path)
    wb = openpyxl.load_workbook(file_path, data_only=True)

    pricing_sheet_name = layout.get("pricing_sheet")
    sla_sheet_name = layout.get("sla_sheet")
    detected_carrier = carrier_hint or layout.get("carrier_name") or "Transportadora"
    detected_table = table_hint or f"Tabela {Path(file_path).stem}"

    if not pricing_sheet_name or pricing_sheet_name not in wb.sheetnames:
        return {"is_valid": False, "issues": ["Aba de tarifas não localizada na planilha."], "rates": []}

    # 1. Extrair tarifas da aba de preços
    pricing_sheet = wb[pricing_sheet_name]
    master_rates = parse_pricing_sheet_matrix(pricing_sheet)

    if not master_rates:
        return {"is_valid": False, "issues": ["Não foi possível extrair matriz tarifária válida."], "rates": []}

    # 2. Extrair cidades e prazos da aba de SLA (se houver)
    cities_map: Dict[str, Dict[str, Dict[str, Any]]] = {}
    if sla_sheet_name and sla_sheet_name in wb.sheetnames:
        sla_sheet = wb[sla_sheet_name]
        cities_map = parse_sla_sheet_cities(sla_sheet)

    # 3. Fusão Híbrida: Regras Mestras por UF + Regras Específicas por Cidade
    final_rates: List[Dict[str, Any]] = []
    mapped_cities = 0

    rates_by_uf_dest: Dict[Tuple[str, str], List[Dict[str, Any]]] = {}
    for r in master_rates:
        rates_by_uf_dest.setdefault((r["uf"], r["dest_type"]), []).append(r)
        final_rates.append(r)  # Regra geral estadual

    for uf, c_dict in cities_map.items():
        for city_name, c_info in c_dict.items():
            mapped_cities += 1
            city_days = c_info["days"]
            tda_val = c_info.get("tda", 0.0)
            city_dest = c_info.get("dest_type", "INTERIOR")

            # Encontrar regra tarifária correspondente (ex: INTERIOR, INTERIOR-2 ou CAPITAL)
            matched_rates = rates_by_uf_dest.get((uf, city_dest))
            if not matched_rates:
                # Fallback: tentar qualquer regra daquela UF
                for (u, _), u_rates in rates_by_uf_dest.items():
                    if u == uf:
                        matched_rates = u_rates
                        break

            if matched_rates:
                for base_r in matched_rates:
                    c_rate = dict(base_r)
                    c_rate["city"] = city_name
                    c_rate["delivery_days"] = city_days
                    tda_note = f" • TDA R${tda_val:.2f}" if tda_val > 0 else ""
                    c_rate["notes"] = f"{city_name}/{uf} ({base_r['dest_type']}){tda_note}".strip()
                    final_rates.append(c_rate)

    summary_msg = f"IA Autônoma: {len(final_rates)} regras compiladas ({mapped_cities} cidades com SLA exato mapeadas)."
    return {
        "is_valid": True,
        "carrier_name": detected_carrier,
        "table_name": detected_table,
        "rates": final_rates,
        "metadata": {
            "origin_city": "Cariacica/ES",
            "cubing_factor": 300.0,
            "tas_fixed": 0.0,
            "notes": summary_msg,
            "mapped_cities": mapped_cities
        },
        "summary": summary_msg
    }

import csv
import io
import unicodedata
from datetime import date
from flask import Blueprint, render_template, request, redirect, url_for, flash, session, jsonify
from werkzeug.utils import secure_filename

try:
    from openpyxl import load_workbook
except Exception:
    load_workbook = None

from database import db
from gemini_service import analyze_import_documents
from routes.helpers import (
    login_required,
    roles_required,
    save_upload,
    audit,
    UPLOAD_DIR
)

import_bp = Blueprint("import", __name__)


def normalize_headers(headers):
    out = []
    for h in headers:
        s = str(h or "").strip().lower()
        s = s.replace("ç", "c").replace("ã", "a").replace("á", "a").replace("é", "e").replace("í", "i").replace("ó", "o").replace("ú", "u")
        out.append(s)
    return out


def parse_chassis_file(file_storage):
    ext = file_storage.filename.rsplit(".", 1)[1].lower()
    data = file_storage.read()
    rows = []
    if ext == "csv":
        text = data.decode("utf-8-sig", errors="replace")
        sample = text[:2048]
        try:
            dialect = csv.Sniffer().sniff(sample, delimiters=",;\t")
        except Exception:
            dialect = csv.excel
            dialect.delimiter = ";"
        reader = csv.reader(io.StringIO(text), dialect)
        all_rows = list(reader)
    elif ext == "xlsx":
        if load_workbook is None:
            raise ValueError("Suporte a XLSX indisponível. Instale openpyxl.")
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        ws = wb.active
        all_rows = [[cell for cell in row] for row in ws.iter_rows(values_only=True)]
    else:
        raise ValueError("Use CSV ou XLSX para a planilha de chassis.")
    if not all_rows:
        return []
    headers = normalize_headers(all_rows[0])
    def idx(candidates):
        for c in candidates:
            if c in headers:
                return headers.index(c)
        return None
    i_model = idx(["modelo", "model", "produto", "product"])
    i_chassis = idx(["chassi", "chassis", "quadro", "frame", "frame no", "frame number", "vin"])
    i_motor = idx(["motor", "motor no", "motor number", "numero do motor", "n motor"])
    i_color = idx(["cor", "color", "colour"])
    if i_model is None or i_chassis is None:
        raise ValueError("A planilha precisa ter pelo menos as colunas MODELO e CHASSI.")
    for raw in all_rows[1:]:
        model = str(raw[i_model] or "").strip() if i_model < len(raw) else ""
        chassis = str(raw[i_chassis] or "").strip() if i_chassis < len(raw) else ""
        motor = str(raw[i_motor] or "").strip() if i_motor is not None and i_motor < len(raw) else ""
        color = str(raw[i_color] or "").strip() if i_color is not None and i_color < len(raw) else ""
        if model and chassis:
            rows.append({"model": model, "chassis": chassis, "motor": motor, "color": color})
    return rows


@import_bp.route("/imports", methods=["GET", "POST"])
@login_required
@roles_required("admin", "support")
def imports():
    if request.method == "POST":
        reference = request.form["reference"].strip()
        invoice_no = request.form.get("invoice_no", "").strip()
        bl_no = request.form.get("bl_no", "").strip()
        supplier_name = request.form.get("supplier_name", "").strip()
        seller_name = request.form.get("seller_name", "").strip()
        arrival_date = request.form.get("arrival_date") or None
        usd_rate = float(request.form.get("usd_rate") or 0)
        invoice_amount_usd = float(request.form.get("invoice_amount_usd") or 0)
        nf_entry = request.form.get("nf_entry", "").strip()
        notes = request.form.get("notes", "").strip()
        try:
            invoice_file = save_upload(request.files.get("invoice_file"), "invoice")
            bl_file = save_upload(request.files.get("bl_file"), "bl")
            nf_entry_file = save_upload(request.files.get("nf_entry_file"), "nfentrada")
            chassis_file_obj = request.files.get("chassis_file")
            chassis_file_name = save_upload(chassis_file_obj, "chassis") if chassis_file_obj and chassis_file_obj.filename else None
        except ValueError as e:
            flash(str(e), "danger")
            return redirect(url_for("imports"))

        with db() as conn:
            cur = conn.execute(
                """INSERT INTO imports(reference,invoice_no,bl_no,supplier_name,seller_name,arrival_date,usd_rate,invoice_amount_usd,nf_entry,invoice_file,bl_file,nf_entry_file,chassis_file,notes,created_by)
                   VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)""",
                (reference, invoice_no, bl_no, supplier_name, seller_name, arrival_date, usd_rate, invoice_amount_usd, nf_entry, invoice_file, bl_file, nf_entry_file, chassis_file_name, notes, session["user_id"]),
            )
            iid = cur.lastrowid

            if chassis_file_name:
                try:
                    class FS:
                        pass
                    obj = FS()
                    obj.filename = chassis_file_name
                    obj.read = lambda: (UPLOAD_DIR / chassis_file_name).read_bytes()
                    rows = parse_chassis_file(obj)
                    inserted = 0
                    for row in rows:
                        existing = conn.execute("SELECT id FROM stock_units WHERE chassis=?", (row["chassis"],)).fetchone()
                        if not existing:
                            prod = conn.execute("SELECT id FROM products WHERE lower(name)=lower(?)", (row["model"],)).fetchone()
                            if not prod:
                                sku_base = "".join(ch for ch in row["model"].upper() if ch.isalnum())[:18] or "PROD"
                                sku = sku_base
                                n = 1
                                while conn.execute("SELECT 1 FROM products WHERE sku=?", (sku,)).fetchone():
                                    n += 1
                                    sku = f"{sku_base}-{n}"
                                cur_p = conn.execute("INSERT INTO products(name,sku,category) VALUES(?,?,?)", (row["model"], sku, "Importado"))
                                product_id = cur_p.lastrowid
                            else:
                                product_id = prod["id"]
                            conn.execute(
                                "INSERT INTO stock_units(chassis,motor_no,product_id,color,import_id,status,received_at) VALUES(?,?,?,?,?,?,?)",
                                (row["chassis"], row["motor"], product_id, row["color"], iid, "unreleased", arrival_date),
                            )
                            inserted += 1
                    flash(f"Importação criada com {inserted} chassis cadastrados.", "success")
                except Exception as ex:
                    flash(f"Importação criada, porém erro ao processar planilha de chassis: {ex}", "warning")
            else:
                flash("Importação criada com sucesso. Carregue a planilha de chassis para liberar o estoque.", "success")
            conn.commit()

        audit("import.created", reference)
        return redirect(url_for("imports"))

    with db() as conn:
        import_rows = conn.execute(
            """
            SELECT i.*, COUNT(DISTINCT st.id) AS chassis_count
            FROM imports i
            LEFT JOIN stock_units st ON st.import_id=i.id
            GROUP BY i.id ORDER BY i.created_at DESC
            """
        ).fetchall()

        result_imports = []
        for row in import_rows:
            imp_dict = dict(row)
            costs = conn.execute(
                "SELECT * FROM import_costs WHERE import_id=? ORDER BY paid_at ASC, id ASC",
                (imp_dict["id"],)
            ).fetchall()

            costs_list = [dict(c) for c in costs]
            supplier_brl = 0.0
            supplier_usd = 0.0
            total_costs_brl = 0.0

            for c in costs_list:
                c_rate = float(c.get("usd_rate") or imp_dict.get("usd_rate") or 1.0)
                amount = float(c.get("amount") or 0.0)
                c_currency = (c.get("currency") or "BRL").upper()

                if c_currency == "USD":
                    amount_usd = amount
                    amount_brl = amount * (c_rate if c_rate > 0 else 1.0)
                else:
                    amount_brl = amount
                    amount_usd = amount / c_rate if c_rate > 0 else 0.0

                c["amount_brl"] = amount_brl
                c["amount_usd"] = amount_usd
                total_costs_brl += amount_brl

                c_type = (c.get("cost_type") or "").strip().lower()
                c_desc = (c.get("description") or "").strip().lower()
                is_supplier_payment = any(kw in c_type or kw in c_desc for kw in ["fornecedor", "sinal", "inicial", "intermedi", "final", "invoice", "china", "chines", "fabricante"])

                if is_supplier_payment or c_currency == "USD":
                    supplier_brl += amount_brl
                    if c_currency == "USD":
                        supplier_usd += amount
                    elif c_rate > 0:
                        supplier_usd += (amount / c_rate)

            imp_dict["costs_list"] = costs_list
            imp_dict["costs_brl"] = total_costs_brl
            imp_dict["supplier_brl"] = supplier_brl
            imp_dict["supplier_usd"] = supplier_usd

            # Cálculo do Câmbio Médio (Dólar Médio dos pagamentos chineses)
            if supplier_usd > 0 and supplier_brl > 0:
                imp_dict["avg_usd_rate"] = round(supplier_brl / supplier_usd, 4)
            elif float(imp_dict.get("invoice_amount_usd") or 0) > 0 and supplier_brl > 0:
                imp_dict["avg_usd_rate"] = round(supplier_brl / float(imp_dict.get("invoice_amount_usd")), 4)
            else:
                imp_dict["avg_usd_rate"] = float(imp_dict.get("usd_rate") or 0.0)

            # Cálculo da Alíquota Real do Dólar da Importação (R$ Total de Despesas / US$ Invoice)
            inv_usd = float(imp_dict.get("invoice_amount_usd") or 0.0)
            if inv_usd > 0 and total_costs_brl > 0:
                imp_dict["landed_aliquota"] = round(total_costs_brl / inv_usd, 4)
            elif supplier_usd > 0 and total_costs_brl > 0:
                imp_dict["landed_aliquota"] = round(total_costs_brl / supplier_usd, 4)
            else:
                imp_dict["landed_aliquota"] = 0.0

            result_imports.append(imp_dict)

    return render_template("imports.html", imports=result_imports)


@import_bp.route("/imports/<int:iid>/edit", methods=["POST"])
@login_required
@roles_required("admin")
def edit_import(iid):
    reference = request.form.get("reference", "").strip()
    invoice_no = request.form.get("invoice_no", "").strip()
    bl_no = request.form.get("bl_no", "").strip()
    supplier_name = request.form.get("supplier_name", "").strip()
    seller_name = request.form.get("seller_name", "").strip()
    nf_entry = request.form.get("nf_entry", "").strip()
    arrival_date = request.form.get("arrival_date") or None
    usd_rate = float(request.form.get("usd_rate") or 0)
    invoice_amount_usd = float(request.form.get("invoice_amount_usd") or 0)
    notes = request.form.get("notes", "").strip()

    with db() as conn:
        imp = conn.execute("SELECT * FROM imports WHERE id=?", (iid,)).fetchone()
        if not imp:
            flash("Importação não encontrada.", "danger")
            return redirect(url_for("imports"))

        try:
            invoice_file = save_upload(request.files.get("invoice_file"), "invoice") or imp["invoice_file"]
            bl_file = save_upload(request.files.get("bl_file"), "bl") or imp["bl_file"]
            nf_entry_file = save_upload(request.files.get("nf_entry_file"), "nfentrada") or imp["nf_entry_file"]
        except ValueError as e:
            flash(str(e), "danger")
            return redirect(url_for("imports"))

        chassis_file = imp["chassis_file"]
        chassis_file_obj = request.files.get("chassis_file")
        if chassis_file_obj and chassis_file_obj.filename:
            try:
                chassis_file = save_upload(chassis_file_obj, "chassis")
                class FS:
                    pass
                obj = FS()
                obj.filename = chassis_file
                obj.read = lambda: (UPLOAD_DIR / chassis_file).read_bytes()
                rows = parse_chassis_file(obj)
                inserted = 0
                for row in rows:
                    existing = conn.execute("SELECT id FROM stock_units WHERE chassis=?", (row["chassis"],)).fetchone()
                    if not existing:
                        prod = conn.execute("SELECT id FROM products WHERE lower(name)=lower(?)", (row["model"],)).fetchone()
                        if not prod:
                            sku_base = "".join(ch for ch in row["model"].upper() if ch.isalnum())[:18] or "PROD"
                            sku = sku_base
                            n = 1
                            while conn.execute("SELECT 1 FROM products WHERE sku=?", (sku,)).fetchone():
                                n += 1
                                sku = f"{sku_base}-{n}"
                            cur_p = conn.execute("INSERT INTO products(name,sku,category) VALUES(?,?,?)", (row["model"], sku, "Importado"))
                            product_id = cur_p.lastrowid
                        else:
                            product_id = prod["id"]
                        conn.execute(
                            "INSERT INTO stock_units(chassis,motor_no,product_id,color,import_id,status,received_at) VALUES(?,?,?,?,?,?,?)",
                            (row["chassis"], row["motor"], product_id, row["color"], iid, "available" if imp["status"] == "released" else "unreleased", arrival_date),
                        )
                        inserted += 1
                flash(f"Planilha de chassis atualizada ({inserted} novos chassis).", "info")
            except Exception as ex:
                flash(f"Erro ao ler planilha de chassis: {ex}", "warning")

        conn.execute(
            """UPDATE imports SET reference=?, invoice_no=?, bl_no=?, supplier_name=?, seller_name=?, nf_entry=?, arrival_date=?, usd_rate=?, invoice_amount_usd=?, invoice_file=?, bl_file=?, nf_entry_file=?, chassis_file=?, notes=?
               WHERE id=?""",
            (reference or imp["reference"], invoice_no, bl_no, supplier_name, seller_name, nf_entry, arrival_date, usd_rate, invoice_amount_usd, invoice_file, bl_file, nf_entry_file, chassis_file, notes, iid)
        )
        conn.commit()
    audit("import.updated", f"import_id={iid}")
    flash("Importação atualizada com sucesso.", "success")
    return redirect(url_for("imports"))


@import_bp.route("/imports/<int:iid>/chassis", methods=["POST"])
@login_required
@roles_required("admin", "stock", "support")
def import_chassis(iid):
    f = request.files.get("chassis_file")
    if not f or not f.filename:
        flash("Selecione uma planilha CSV ou XLSX.", "danger")
        return redirect(url_for("imports"))
    try:
        filename = save_upload(f, "chassis")
        class FS:
            pass
        obj = FS()
        obj.filename = filename
        obj.read = lambda: (UPLOAD_DIR / filename).read_bytes()
        rows = parse_chassis_file(obj)
    except Exception as e:
        flash(f"Não foi possível importar a planilha: {e}", "danger")
        return redirect(url_for("imports"))
    inserted = 0
    duplicates = []
    with db() as conn:
        imp = conn.execute("SELECT * FROM imports WHERE id=?", (iid,)).fetchone()
        if not imp:
            flash("Importação não encontrada.", "danger")
            return redirect(url_for("imports"))
        for row in rows:
            existing = conn.execute("SELECT id FROM stock_units WHERE chassis=?", (row["chassis"],)).fetchone()
            if existing:
                duplicates.append(row["chassis"])
                continue
            prod = conn.execute("SELECT id FROM products WHERE lower(name)=lower(?)", (row["model"],)).fetchone()
            if not prod:
                sku_base = "".join(ch for ch in row["model"].upper() if ch.isalnum())[:18] or "PROD"
                sku = sku_base
                n = 1
                while conn.execute("SELECT 1 FROM products WHERE sku=?", (sku,)).fetchone():
                    n += 1
                    sku = f"{sku_base}-{n}"
                cur = conn.execute("INSERT INTO products(name,sku,category) VALUES(?,?,?)", (row["model"], sku, "Importado"))
                product_id = cur.lastrowid
            else:
                product_id = prod["id"]
            conn.execute(
                "INSERT INTO stock_units(chassis,motor_no,product_id,color,import_id,status,received_at) VALUES(?,?,?,?,?,?,?)",
                (row["chassis"], row["motor"], product_id, row["color"], iid, "available" if imp["status"] == "released" else "unreleased", imp["arrival_date"]),
            )
            inserted += 1
        conn.execute("UPDATE imports SET chassis_file=? WHERE id=?", (filename, iid))
        conn.commit()
    audit("import.chassis", f"import_id={iid}; inserted={inserted}; duplicates={len(duplicates)}")
    msg = f"{inserted} chassis importados."
    if duplicates:
        msg += f" {len(duplicates)} duplicados foram bloqueados."
    flash(msg, "success" if inserted else "warning")
    return redirect(url_for("imports"))


@import_bp.route("/imports/<int:iid>/cost", methods=["POST"])
@login_required
@roles_required("admin", "support")
def add_import_cost(iid):
    try:
        receipt = save_upload(request.files.get("receipt_file"), "importcost")
    except ValueError as e:
        flash(str(e), "danger")
        return redirect(url_for("imports"))
    with db() as conn:
        conn.execute(
            "INSERT INTO import_costs(import_id,cost_type,description,amount,currency,usd_rate,paid_at,receipt_file) VALUES(?,?,?,?,?,?,?,?)",
            (
                iid,
                request.form.get("cost_type", "Pagamento Extra / Outros"),
                request.form.get("description", "").strip(),
                float(request.form.get("amount") or 0),
                request.form.get("currency", "BRL"),
                float(request.form.get("usd_rate") or 0),
                request.form.get("paid_at") or date.today().isoformat(),
                receipt,
            ),
        )
        conn.commit()
    audit("import.cost", f"import_id={iid}")
    flash("Comprovante/Custo adicionado à importação.", "success")
    return redirect(url_for("imports"))


@import_bp.route("/imports/cost/<int:cid>/delete", methods=["POST"])
@login_required
@roles_required("admin")
def delete_import_cost(cid):
    with db() as conn:
        cost = conn.execute("SELECT import_id FROM import_costs WHERE id=?", (cid,)).fetchone()
        if cost:
            conn.execute("DELETE FROM import_costs WHERE id=?", (cid,))
            conn.commit()
            audit("import.cost_deleted", f"cost_id={cid}")
            flash("Comprovante/Custo removido com sucesso.", "success")
    return redirect(url_for("imports"))


@import_bp.route("/imports/<int:iid>/release", methods=["POST"])
@login_required
@roles_required("admin", "support")
def release_import(iid):
    with db() as conn:
        imp = conn.execute("SELECT * FROM imports WHERE id=?", (iid,)).fetchone()
        count = conn.execute("SELECT COUNT(*) c FROM stock_units WHERE import_id=?", (iid,)).fetchone()["c"]
        if not imp:
            flash("Importação não encontrada.", "danger")
        elif count == 0:
            flash("BLOQUEIO DE SEGURANÇA: Não é possível liberar a importação para venda sem cadastrar os chassis da remessa. Carregue a planilha de chassis primeiro.", "danger")
        elif not imp["invoice_no"] or not imp["bl_no"]:
            flash("Para liberar a importação, é necessário informar Invoice e BL.", "danger")
        else:
            conn.execute("UPDATE imports SET status='released' WHERE id=?", (iid,))
            conn.execute("UPDATE stock_units SET status='available' WHERE import_id=? AND status='unreleased'", (iid,))
            conn.commit()
            audit("import.released", f"import_id={iid}")
            flash("Estoque desta importação foi liberado com sucesso para venda.", "success")
    return redirect(url_for("imports"))


@import_bp.route("/api/imports/analyze-docs", methods=["POST"])
@login_required
@roles_required("admin", "support")
def api_analyze_import_docs():
    file_objs = []
    for key in ["invoice_file", "bl_file", "nf_entry_file", "chassis_file"]:
        f = request.files.get(key)
        if f and f.filename:
            content = f.read()
            ext = f.filename.rsplit(".", 1)[-1].lower()
            mime = "application/pdf" if ext == "pdf" else ("text/csv" if ext == "csv" else f"image/{ext if ext != 'jpg' else 'jpeg'}")
            file_objs.append({
                "bytes": content,
                "mime_type": mime,
                "filename": f.filename
            })

    if not file_objs:
        return jsonify({"success": False, "message": "Nenhum arquivo enviado para análise da IA."})

    res = analyze_import_documents(file_objs)
    return jsonify(res)

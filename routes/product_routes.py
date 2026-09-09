"""
M-One Products & Pricing Blueprint (routes/product_routes.py)
Rotas de Catálogo de Produtos, Tabela de Preços, Importação de Planilhas e Sincronização Bling/Webhooks.
"""

from __future__ import annotations

import csv
import io
import os
import unicodedata
import urllib.request

from flask import Blueprint, flash, jsonify, redirect, render_template, request, url_for

try:
    from openpyxl import load_workbook
except Exception:
    load_workbook = None

from database import db
from routes.helpers import (
    audit,
    login_required,
    roles_required,
    normalize_headers,
    find_header_and_data_rows,
)

product_bp = Blueprint("products", __name__)


def ensure_product_columns():
    try:
        with db() as conn:
            cols = [("fob_price_usd", "REAL DEFAULT 0"), ("aliquota_rate", "REAL DEFAULT 0"), ("installment_12x", "REAL DEFAULT 0"), ("installment_18x", "REAL DEFAULT 0"), ("bling_id", "TEXT"), ("bling_stock", "INTEGER DEFAULT 0"), ("bling_updated_at", "TIMESTAMP")]
            for col, col_type in cols:
                try:
                    conn.execute(f"ALTER TABLE products ADD COLUMN {col} {col_type}")
                except Exception:
                    pass
            conn.commit()
    except Exception:
        pass


def clean_product_name(raw_name: str) -> str:
    return (raw_name or "").strip()


def calc_product_pricing(fob, aliq, cost, retail, inst12, inst18):
    if fob > 0 and aliq > 0:
        cost = round(fob * aliq, 2)
    elif cost > 0 and fob > 0 and aliq == 0:
        aliq = round(cost / fob, 4)
    elif cost > 0 and aliq > 0 and fob == 0:
        fob = round(cost / aliq, 2)
    if retail > 0:
        if inst12 == 0:
            inst12 = round((retail * 1.1013216) / 12, 2)
        if inst18 == 0:
            inst18 = round((retail * 1.1437722) / 18, 2)
    return fob, aliq, cost, inst12, inst18


def parse_products_rows(data_bytes=None, text_content=None, filename="sheet.csv"):
    rows = []
    all_rows = []
    if text_content:
        text = text_content
        sample = text[:2048]
        try:
            dialect = csv.Sniffer().sniff(sample, delimiters=",;\t")
        except Exception:
            dialect = csv.excel
            dialect.delimiter = ";"
        reader = csv.reader(io.StringIO(text), dialect)
        all_rows = list(reader)
    elif data_bytes:
        ext = filename.rsplit(".", 1)[-1].lower() if "." in filename else "csv"
        if ext == "xlsx":
            if load_workbook is None:
                raise ValueError("Suporte a XLSX indisponível. Instale openpyxl.")
            wb = load_workbook(io.BytesIO(data_bytes), read_only=True, data_only=True)
            ws = wb.active
            all_rows = [[cell for cell in row] for row in ws.iter_rows(values_only=True)]
        else:
            text = data_bytes.decode("utf-8-sig", errors="replace")
            sample = text[:2048]
            try:
                dialect = csv.Sniffer().sniff(sample, delimiters=",;\t")
            except Exception:
                dialect = csv.excel
                dialect.delimiter = ";"
            reader = csv.reader(io.StringIO(text), dialect)
            all_rows = list(reader)

    if not all_rows:
        return []

    h_idx, headers, data_rows = find_header_and_data_rows(all_rows)

    def idx(candidates):
        for c in candidates:
            if c in headers:
                return headers.index(c)
        return None

    i_name = idx(["produto", "product", "modelo", "model", "nome", "name", "descricao", "description"])
    i_sku = idx(["sku", "codigo", "code"])
    i_category = idx(["categoria", "category", "tipo", "type"])
    i_fob = idx(["fob", "fob (usd)", "fob usd", "preco fob", "usd fob", "valor fob"])
    i_aliquota = idx(["aliquota", "aliquota real", "aliquota (r$/usd)", "taxa aliquota", "aliquota r$"])
    i_cost = idx(["custo", "unit_cost", "cost", "custo unitario", "valor custo", "custo galpao"])
    i_wholesale = idx(["atacado", "wholesale", "wholesale_price", "preco atacado", "valor atacado"])
    i_retail = idx(["varejo", "retail", "retail_price", "preco varejo", "valor varejo", "preco", "price"])
    i_inst_12x = idx(["12x varejo", "12x", "12x (parcela)", "parcela 12x", "12x varejo (parcela)"])
    i_inst_18x = idx(["18x varejo", "18x", "18x (parcela)", "parcela 18x", "18x varejo (parcela)"])
    i_promo = idx(["promo", "promo_eligible", "elegivel", "promocional"])

    if i_name is None and i_sku is None:
        raise ValueError("A planilha precisa conter ao menos a coluna 'PRODUTO', 'MODELO' ou 'SKU'.")

    def parse_float(val):
        if val is None:
            return 0.0
        s = str(val).strip().replace("R$", "").replace("$", "").replace(" ", "").replace(".", "").replace(",", ".")
        try:
            return float(s)
        except Exception:
            try:
                return float(str(val).strip().replace(",", "."))
            except Exception:
                return 0.0

    for raw in data_rows:
        name = str(raw[i_name] or "").strip() if i_name is not None and i_name < len(raw) else ""
        sku = str(raw[i_sku] or "").strip() if i_sku is not None and i_sku < len(raw) else ""
        category = str(raw[i_category] or "").strip() if i_category is not None and i_category < len(raw) else ""
        fob_raw = parse_float(raw[i_fob]) if i_fob is not None and i_fob < len(raw) else 0.0
        aliq_raw = parse_float(raw[i_aliquota]) if i_aliquota is not None and i_aliquota < len(raw) else 0.0
        cost_raw = parse_float(raw[i_cost]) if i_cost is not None and i_cost < len(raw) else 0.0
        wholesale_price = parse_float(raw[i_wholesale]) if i_wholesale is not None and i_wholesale < len(raw) else 0.0
        retail_price = parse_float(raw[i_retail]) if i_retail is not None and i_retail < len(raw) else 0.0
        inst_12x = parse_float(raw[i_inst_12x]) if i_inst_12x is not None and i_inst_12x < len(raw) else 0.0
        inst_18x = parse_float(raw[i_inst_18x]) if i_inst_18x is not None and i_inst_18x < len(raw) else 0.0

        fob_price_usd, aliquota_rate, unit_cost, installment_12x, installment_18x = calc_product_pricing(
            fob_raw, aliq_raw, cost_raw, retail_price, inst_12x, inst_18x
        )

        promo_val = str(raw[i_promo] or "").strip().lower() if i_promo is not None and i_promo < len(raw) else "1"
        promo_eligible = True if promo_val in ["1", "true", "sim", "s", "elegivel", "yes"] else False

        if name or sku:
            rows.append({
                "name": clean_product_name(name) if name else sku,
                "sku": sku or None,
                "category": category or None,
                "fob_price_usd": fob_price_usd,
                "aliquota_rate": aliquota_rate,
                "unit_cost": unit_cost,
                "wholesale_price": wholesale_price,
                "retail_price": retail_price,
                "installment_12x": installment_12x,
                "installment_18x": installment_18x,
                "promo_eligible": promo_eligible
            })
    return rows


@product_bp.route("/products", methods=["GET", "POST"])
@login_required
@roles_required("admin", "finance", "stock", "support")
def products():
    ensure_product_columns()

    if request.method == "POST":
        name = request.form["name"].strip()
        sku = request.form.get("sku", "").strip() or None
        category = request.form.get("category", "").strip()
        retail_price = float(request.form.get("retail_price") or 0)
        wholesale_price = float(request.form.get("wholesale_price") or 0)
        fob_price_usd, aliquota_rate, unit_cost, installment_12x, installment_18x = calc_product_pricing(
            float(request.form.get("fob_price_usd") or 0),
            float(request.form.get("aliquota_rate") or 0),
            float(request.form.get("unit_cost") or 0),
            retail_price,
            float(request.form.get("installment_12x") or 0),
            float(request.form.get("installment_18x") or 0)
        )
        promo_eligible = True if request.form.get("promo_eligible") else False

        with db() as conn:
            conn.execute(
                """INSERT INTO products(name,sku,category,fob_price_usd,aliquota_rate,unit_cost,retail_price,wholesale_price,installment_12x,installment_18x,promo_eligible)
                   VALUES(%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s) RETURNING id""",
                (name, sku, category, fob_price_usd, aliquota_rate, unit_cost, retail_price, wholesale_price, installment_12x, installment_18x, promo_eligible),
            )
            conn.commit()
        audit("product.created", name)
        flash("Produto cadastrado.", "success")
        return redirect(url_for("products"))

    with db() as conn:
        rows = conn.execute(
            """
            SELECT p.*,
                   COALESCE(p.bling_stock, 0) as bling_stock,
                   p.bling_updated_at,
                   COALESCE(st.local_available, 0) as local_available,
                   COALESCE(st.sold, 0) as sold
            FROM products p
            LEFT JOIN (
                SELECT product_id,
                       SUM(CASE WHEN status='available' THEN 1 ELSE 0 END) as local_available,
                       SUM(CASE WHEN status='sold' THEN 1 ELSE 0 END) as sold
                FROM stock_units
                GROUP BY product_id
            ) st ON st.product_id = p.id
            ORDER BY p.name
            """
        ).fetchall()

        last_sync = None
        for r in rows:
            if r["bling_updated_at"]:
                last_sync = r["bling_updated_at"]
                break

    return render_template("products.html", products=rows, last_sync=last_sync)


@product_bp.route("/products/<int:pid>/edit", methods=["POST"])
@login_required
@roles_required("admin", "finance", "support")
def edit_product(pid: int):
    retail_price = float(request.form.get("retail_price") or 0)
    wholesale_price = float(request.form.get("wholesale_price") or 0)
    fob_price_usd, aliquota_rate, unit_cost, installment_12x, installment_18x = calc_product_pricing(
        float(request.form.get("fob_price_usd") or 0),
        float(request.form.get("aliquota_rate") or 0),
        float(request.form.get("unit_cost") or 0),
        retail_price,
        float(request.form.get("installment_12x") or 0),
        float(request.form.get("installment_18x") or 0)
    )

    fields = (
        request.form.get("name", "").strip(),
        request.form.get("sku", "").strip() or None,
        request.form.get("category", "").strip(),
        fob_price_usd,
        aliquota_rate,
        unit_cost,
        retail_price,
        wholesale_price,
        installment_12x,
        installment_18x,
        1 if request.form.get("promo_eligible") else 0,
        pid,
    )
    with db() as conn:
        conn.execute(
            """UPDATE products SET name=%s,sku=%s,category=%s,fob_price_usd=%s,aliquota_rate=%s,unit_cost=%s,retail_price=%s,wholesale_price=%s,installment_12x=%s,installment_18x=%s,promo_eligible=%s
               WHERE id=%s""",
            fields,
        )
        conn.commit()
    audit("product.updated", f"product_id={pid}")
    flash("Produto atualizado.", "success")
    return redirect(url_for("products"))


@product_bp.route("/products/import", methods=["POST"])
@login_required
@roles_required("admin", "finance", "stock")
def import_products():
    sheets_url = request.form.get("sheets_url", "").strip()
    file_storage = request.files.get("products_file")

    parsed_rows = []

    if sheets_url:
        export_url = sheets_url
        if "docs.google.com/spreadsheets" in export_url and "/export" not in export_url:
            export_url = export_url.split("/edit")[0].rstrip("/") + "/export?format=csv"
            gid = None
            if "#gid=" in sheets_url:
                gid = sheets_url.split("#gid=")[1].split("&")[0]
            elif "gid=" in sheets_url:
                gid = sheets_url.split("gid=")[1].split("&")[0]
            if gid:
                export_url += f"&gid={gid}"

        try:
            req = urllib.request.Request(export_url, headers={"User-Agent": "Mozilla/5.0"})
            with urllib.request.urlopen(req, timeout=15) as resp:
                csv_text = resp.read().decode("utf-8-sig", errors="replace")
                parsed_rows = parse_products_rows(text_content=csv_text)
        except Exception as e:
            flash(f"Falha ao carregar a planilha do Google Sheets pelo link. Verifique se a permissão do link está como 'Qualquer pessoa com o link pode ver'. Erro: {str(e)}", "danger")
            return redirect(url_for("products"))

    elif file_storage and file_storage.filename:
        data = file_storage.read()
        try:
            parsed_rows = parse_products_rows(data_bytes=data, filename=file_storage.filename)
        except Exception as e:
            flash(f"Erro ao processar arquivo: {str(e)}", "danger")
            return redirect(url_for("products"))
    else:
        flash("Selecione um arquivo CSV/XLSX ou informe a URL do Google Sheets.", "danger")
        return redirect(url_for("products"))

    if not parsed_rows:
        flash("Nenhum produto válido encontrado na planilha.", "warning")
        return redirect(url_for("products"))

    created_count = 0
    updated_count = 0

    with db() as conn:
        for r in parsed_rows:
            existing = None
            if r["sku"]:
                existing = conn.execute("SELECT id FROM products WHERE lower(sku)=lower(%s)", (r["sku"],)).fetchone()
            if not existing and r["name"]:
                existing = conn.execute("SELECT id FROM products WHERE lower(name)=lower(%s)", (r["name"],)).fetchone()

            if existing:
                conn.execute(
                    """UPDATE products SET name=%s, sku=COALESCE(%s, sku), category=COALESCE(%s, category),
                       fob_price_usd=%s, aliquota_rate=%s, unit_cost=%s, wholesale_price=%s, retail_price=%s, installment_12x=%s, installment_18x=%s, promo_eligible=%s WHERE id=%s""",
                    (r["name"], r["sku"], r["category"], r.get("fob_price_usd", 0), r.get("aliquota_rate", 0), r["unit_cost"], r["wholesale_price"], r["retail_price"], r.get("installment_12x", 0), r.get("installment_18x", 0), r["promo_eligible"], existing["id"])
                )
                updated_count += 1
            else:
                conn.execute(
                    """INSERT INTO products (name, sku, category, fob_price_usd, aliquota_rate, unit_cost, wholesale_price, retail_price, installment_12x, installment_18x, promo_eligible)
                       VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s)""",
                    (r["name"], r["sku"], r["category"], r.get("fob_price_usd", 0), r.get("aliquota_rate", 0), r["unit_cost"], r["wholesale_price"], r["retail_price"], r.get("installment_12x", 0), r.get("installment_18x", 0), r["promo_eligible"])
                )
                created_count += 1
        conn.commit()

    audit("products.imported", f"created={created_count}; updated={updated_count}")
    flash(f"Planilha de produtos processada! {created_count} novos produtos cadastrados, {updated_count} atualizados.", "success")
    return redirect(url_for("products"))


@product_bp.route("/api/sync-prices", methods=["POST"])
@product_bp.route("/sync-prices", methods=["POST"])
def api_sync_prices():
    token = request.args.get("token") or request.headers.get("X-Sync-Token")
    expected_token = os.environ.get("SYNC_TOKEN") or "maj-m-one-sync-secret-2026"
    if not token or token != expected_token:
        return jsonify({"ok": False, "error": "Token de autenticação inválido"}), 401

    payload = request.get_json(silent=True)
    all_rows = []

    if isinstance(payload, list):
        all_rows = payload
    elif isinstance(payload, dict):
        all_rows = payload.get("data") or payload.get("rows") or payload.get("products") or []

    if not all_rows and request.data:
        text = request.data.decode("utf-8-sig", errors="replace")
        try:
            dialect = csv.Sniffer().sniff(text[:2048], delimiters=",;\t")
        except Exception:
            dialect = csv.excel
            dialect.delimiter = ";"
        reader = csv.reader(io.StringIO(text), dialect)
        all_rows = list(reader)

    if not all_rows:
        return jsonify({"ok": False, "error": "Nenhum dado enviado"}), 400

    parsed_rows = []
    if isinstance(all_rows[0], list):
        h_idx, headers, data_rows = find_header_and_data_rows(all_rows)

        def idx(candidates):
            for c in candidates:
                if c in headers:
                    return headers.index(c)
            return None

        i_name = idx(["produto", "product", "modelo", "model", "nome", "name", "descricao", "description"])
        i_sku = idx(["sku", "codigo", "code"])
        i_category = idx(["categoria", "category", "tipo", "type"])
        i_cost = idx(["custo", "unit_cost", "cost", "custo unitario", "valor custo"])
        i_wholesale = idx(["atacado", "wholesale", "wholesale_price", "preco atacado", "valor atacado"])
        i_retail = idx(["varejo", "retail", "retail_price", "preco varejo", "valor varejo", "preco", "price"])
        i_inst_12x = idx(["12x varejo", "12x", "12x (parcela)", "parcela 12x", "12x varejo (parcela)"])
        i_inst_18x = idx(["18x varejo", "18x", "18x (parcela)", "parcela 18x", "18x varejo (parcela)"])
        i_promo = idx(["promo", "promo_eligible", "elegivel", "promocional"])

        def parse_float(val):
            if val is None:
                return 0.0
            s = str(val).strip().replace("R$", "").replace(" ", "").replace(".", "").replace(",", ".")
            try:
                return float(s)
            except Exception:
                try:
                    return float(str(val).strip().replace(",", "."))
                except Exception:
                    return 0.0

        for raw in data_rows:
            name = str(raw[i_name] or "").strip() if i_name is not None and i_name < len(raw) else ""
            sku = str(raw[i_sku] or "").strip() if i_sku is not None and i_sku < len(raw) else ""
            category = str(raw[i_category] or "").strip() if i_category is not None and i_category < len(raw) else ""
            unit_cost = parse_float(raw[i_cost]) if i_cost is not None and i_cost < len(raw) else 0.0
            wholesale_price = parse_float(raw[i_wholesale]) if i_wholesale is not None and i_wholesale < len(raw) else 0.0
            retail_price = parse_float(raw[i_retail]) if i_retail is not None and i_retail < len(raw) else 0.0
            installment_12x = parse_float(raw[i_inst_12x]) if i_inst_12x is not None and i_inst_12x < len(raw) else 0.0
            installment_18x = parse_float(raw[i_inst_18x]) if i_inst_18x is not None and i_inst_18x < len(raw) else 0.0

            promo_val = str(raw[i_promo] or "").strip().lower() if i_promo is not None and i_promo < len(raw) else "1"
            promo_eligible = True if promo_val in ["1", "true", "sim", "s", "elegivel", "yes"] else False

            if name or sku:
                parsed_rows.append({
                    "name": clean_product_name(name) if name else sku,
                    "sku": sku or None,
                    "category": category or None,
                    "unit_cost": unit_cost,
                    "wholesale_price": wholesale_price,
                    "retail_price": retail_price,
                    "installment_12x": installment_12x,
                    "installment_18x": installment_18x,
                    "promo_eligible": promo_eligible
                })
    elif isinstance(all_rows[0], dict):
        for item in all_rows:
            name = str(item.get("name") or item.get("produto") or item.get("modelo") or "").strip()
            sku = str(item.get("sku") or item.get("codigo") or "").strip() or None
            category = str(item.get("category") or item.get("categoria") or "").strip() or None
            unit_cost = float(item.get("unit_cost") or item.get("custo") or 0)
            wholesale_price = float(item.get("wholesale_price") or item.get("atacado") or 0)
            retail_price = float(item.get("retail_price") or item.get("varejo") or 0)
            installment_12x = float(item.get("installment_12x") or item.get("12x") or 0)
            installment_18x = float(item.get("installment_18x") or item.get("18x") or 0)
            promo_eligible = bool(item.get("promo_eligible", True))
            if name or sku:
                parsed_rows.append({
                    "name": clean_product_name(name) if name else sku,
                    "sku": sku,
                    "category": category,
                    "unit_cost": unit_cost,
                    "wholesale_price": wholesale_price,
                    "retail_price": retail_price,
                    "installment_12x": installment_12x,
                    "installment_18x": installment_18x,
                    "promo_eligible": promo_eligible
                })

    created_count = 0
    updated_count = 0

    with db() as conn:
        for r in parsed_rows:
            existing = None
            if r["sku"]:
                existing = conn.execute("SELECT id FROM products WHERE lower(sku)=lower(%s)", (r["sku"],)).fetchone()
            if not existing and r["name"]:
                existing = conn.execute("SELECT id FROM products WHERE lower(name)=lower(%s)", (r["name"],)).fetchone()

            if existing:
                conn.execute(
                    """UPDATE products SET name=%s, sku=COALESCE(%s, sku), category=COALESCE(%s, category),
                       unit_cost=%s, wholesale_price=%s, retail_price=%s, installment_12x=%s, installment_18x=%s, promo_eligible=%s WHERE id=%s""",
                    (r["name"], r["sku"], r["category"], r["unit_cost"], r["wholesale_price"], r["retail_price"], r.get("installment_12x", 0), r.get("installment_18x", 0), r["promo_eligible"], existing["id"])
                )
                updated_count += 1
            else:
                conn.execute(
                    """INSERT INTO products (name, sku, category, unit_cost, wholesale_price, retail_price, installment_12x, installment_18x, promo_eligible)
                       VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s)""",
                    (r["name"], r["sku"], r["category"], r["unit_cost"], r["wholesale_price"], r["retail_price"], r.get("installment_12x", 0), r.get("installment_18x", 0), r["promo_eligible"])
                )
                created_count += 1
        conn.commit()

    audit("products.sync_webhook", f"created={created_count}; updated={updated_count}")
    return jsonify({
        "ok": True,
        "message": f"Sincronização concluída com sucesso. {created_count} criados, {updated_count} atualizados.",
        "created": created_count,
        "updated": updated_count
    })

